"""Atualiza tudo com UM comando, na ordem certa: teste antes do treino.

    python atualizar.py caminho/da/planilha.xlsx

O .xlsx é a planilha "Bola de Ouro" exportada do Google Sheets (todas as abas; só 2026 e Jogos são usadas).
O que acontece:
  1. extrai as abas 2026 e Jogos para dados/ (bola_de_ouro_2026_export.txt e jogos.csv)
  2. TESTE PROSPECTIVO: o modelo antigo (dados/data.json) prevê os sábados novos     -> dados/historico_previsoes.csv
  3. posições: atletas novos recebem posição proposta (as já aceitas não mudam)
  4. RETREINO: ratings, goleiros e pares; atualiza dados/data.json (regras e rulesVer são mantidos)
  5. teste retroativo dos últimos 8 sábados                                          -> dados/historico_modelo.csv
  6. aba Jogos: força prevista x resultado de cada jogo e confronto                  -> dados/jogos_confrontos.csv, historico_jogos.csv
  7. reconstrói index.html
Se algo falhar, os arquivos de dados voltam ao que eram antes. Não faz commit nem publica.
"""
import sys, os, csv, json, shutil, subprocess, datetime, warnings
warnings.filterwarnings('ignore')
HERE = os.path.dirname(os.path.abspath(__file__)); os.chdir(HERE)
DADOS = os.path.join('..', 'dados'); APP = os.path.join('..', 'app')
MES = {1: 'Jan', 2: 'Feb', 3: 'Mar', 4: 'Apr', 5: 'May', 6: 'Jun', 7: 'Jul', 8: 'Aug', 9: 'Sep', 10: 'Oct', 11: 'Nov', 12: 'Dec'}
ARQ = ['bola_de_ouro_2026_export.txt', 'jogos.csv', 'data.json', 'posicoes_propostas.csv', 'historico_modelo.csv', 'historico_previsoes.csv', 'historico_jogos.csv', 'jogos_confrontos.csv']

def passo(n, txt): print(f"\n[{n}/7] {txt}", flush=True)
def roda(script, *args, mostrar=False, cwd=None):
    r = subprocess.run([sys.executable, script, *args], capture_output=True, text=True, cwd=cwd or HERE)
    if r.returncode != 0:
        raise RuntimeError(f"{script} falhou:\n{r.stdout[-1500:]}\n{r.stderr[-1500:]}")
    if mostrar: print(r.stdout.rstrip())
    return r.stdout

def extrair(xlsx):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    for aba in ('2026', 'Jogos'):
        if aba not in wb.sheetnames: raise RuntimeError(f"A planilha não tem a aba '{aba}'. Abas: {wb.sheetnames}")
    def cel(v):
        if v is None: return ''
        if isinstance(v, datetime.datetime): return f"{v.day:02d}/{MES[v.month]}/{v.year}"
        if isinstance(v, float) and v.is_integer(): return int(v)
        return v
    linhas = [[cel(v) for v in r[:17]] for r in wb['2026'].iter_rows(values_only=True)]
    linhas = [l for l in linhas if any(c != '' for c in l)]
    with open(os.path.join(DADOS, 'bola_de_ouro_2026_export.txt'), 'w', newline='', encoding='utf-8') as f:
        f.write('## Sheet name: 2026\n'); csv.writer(f, lineterminator='\n').writerows(linhas)
    with open(os.path.join(DADOS, 'jogos.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f, lineterminator='\n'); w.writerow(['data', 'jogo', 'time_a', 'gols_a', 'gols_b', 'time_b'])
        n = 0
        for r in list(wb['Jogos'].iter_rows(values_only=True))[1:]:
            if r[1] is None: continue
            w.writerow([r[1].strftime('%Y-%m-%d'), cel(r[2]), r[3] or '', cel(r[4]), cel(r[5]), r[6] or '']); n += 1
    print(f"   aba 2026: {len(linhas)-1} linhas | aba Jogos: {n} jogos")

def main():
    if len(sys.argv) < 2 or not os.path.exists(sys.argv[1]):
        print(__doc__); sys.exit(1)
    xlsx = os.path.abspath(sys.argv[1])
    backup = {a: open(os.path.join(DADOS, a), 'rb').read() for a in ARQ if os.path.exists(os.path.join(DADOS, a))}
    velho = json.load(open(os.path.join(DADOS, 'data.json'), encoding='utf-8'))
    try:
        passo(1, 'Lendo a planilha'); extrair(xlsx)
        passo(2, 'Teste prospectivo (modelo antigo x sábados novos)'); print(roda('avaliar_novos.py').rstrip())
        passo(3, 'Posições de atletas novos')
        pos_ant = os.path.join(DADOS, 'posicoes_propostas.csv'); import pandas as pd
        antigo = pd.read_csv(pos_ant)
        roda('classify.py'); roda('final.py')          # final.py reescreve o CSV inteiro
        novo = pd.read_csv(pos_ant)
        add = novo[~novo.Atleta.isin(antigo.Atleta)]
        pd.concat([antigo, add], ignore_index=True).to_csv(pos_ant, index=False, encoding='utf-8-sig')
        print("   atletas novos:", ', '.join(f"{r.Atleta} ({r['Posição proposta']}, confiança {r['Confiança']})" for _, r in add.iterrows()) or 'nenhum')
        passo(4, 'Retreino')
        for s in ('final_fit.py', 'gk.py', 'export.py'): roda(s)
        novo_json = json.load(open(os.path.join(DADOS, 'data_modelo.json'), encoding='utf-8')); os.remove(os.path.join(DADOS, 'data_modelo.json'))
        assert len(novo_json['roster']) >= len(velho['roster']), 'o novo elenco ficou menor que o antigo'
        assert novo_json['meta']['validos'] >= velho['meta']['validos'], 'menos sábados válidos que antes'
        d = dict(velho); d.update({k: novo_json[k] for k in ('roster', 'pairs', 'pairstat', 'meta')})
        txt = json.dumps(d, ensure_ascii=False, allow_nan=False)
        open(os.path.join(DADOS, 'data.json'), 'w', encoding='utf-8').write(txt)
        m0, m1 = velho['meta'], d['meta']
        print(f"   base até {m0['ate']} -> {m1['ate']} | sábados válidos {m0['validos']} -> {m1['validos']} | R² (um dia fora) {m0['r2']} -> {m1['r2']} | atletas {len(velho['roster'])} -> {len(d['roster'])}")
        passo(5, 'Teste retroativo (últimos 8 sábados)'); print('\n'.join(roda('avaliar_sabados.py', '8').splitlines()[-4:-1]))
        passo(6, 'Aba Jogos'); print(roda('jogos.py').rstrip())
        passo(7, 'Reconstruindo index.html'); roda('build.py', cwd=os.path.join(HERE, APP)); print('   index.html atualizado')
    except Exception as e:
        for a, b in backup.items(): open(os.path.join(DADOS, a), 'wb').write(b)
        print(f"\nERRO: {e}\nOs arquivos de dados foram restaurados ao estado anterior.")
        sys.exit(1)
    finally:
        for lixo in ('__pycache__',): shutil.rmtree(lixo, ignore_errors=True)
        for p in os.listdir('.'):
            if p.endswith('.pkl'): os.remove(p)
    print("\nPronto. Revise com `git status` e `git diff --stat`; depois faça o commit de sorteio_am/.")

main()
